import io
import json
import os
import re
from typing import Any

import httpx
from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from pydantic import BaseModel, Field
from dotenv import load_dotenv

load_dotenv()

router = APIRouter(prefix="/api/resume-parser", tags=["Resume Parser"])

MAX_FILE_BYTES = 12 * 1024 * 1024
CATEGORIES = {
    "COD", "DSA", "OOD", "APTI", "COMM", "AI",
    "CLOUD", "SQL", "SWE", "SYSD", "NETW", "OS", "OTHER"
}


class Candidate(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None


class Education(BaseModel):
    degree: str | None = None
    institution: str | None = None
    graduation_year: str | int | None = None


class Project(BaseModel):
    name: str
    description: str | None = None
    technologies: list[str] = Field(default_factory=list)


class Skill(BaseModel):
    name: str
    category: str
    evidence: str
    validation_status: str = "needs_review"


class ResumeResult(BaseModel):
    candidate: Candidate
    education: list[Education] = Field(default_factory=list)
    projects: list[Project] = Field(default_factory=list)
    skills: list[Skill] = Field(default_factory=list)
    extraction_warnings: list[str] = Field(default_factory=list)


async def extract_text(filename: str, content: bytes) -> str:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""

    try:
        if suffix == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            return "\n".join(
                page.extract_text() or "" for page in reader.pages
            ).strip()

        if suffix == "docx":
            from docx import Document

            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]

            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))

            return "\n".join(parts).strip()

    except ImportError as exc:
        raise HTTPException(
            status_code=503,
            detail=f"Missing extraction dependency: {exc.name}. Install backend requirements."
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=422,
            detail=f"Could not read this document: {exc}"
        ) from exc

    raise HTTPException(status_code=415, detail="Upload a PDF or DOCX file.")


def validate_result(raw: Any, resume_text: str) -> dict:
    if not isinstance(raw, dict):
        raise HTTPException(status_code=502, detail="Gemini returned an invalid response.")

    skills = []
    raw_skills = raw.get("skills", [])

    if not isinstance(raw_skills, list):
        raw_skills = []

    for item in raw_skills:
        if not isinstance(item, dict):
            continue

        name = str(item.get("name", "")).strip()
        if not name:
            continue

        category = str(item.get("category", "OTHER")).upper()
        if category not in CATEGORIES:
            category = "OTHER"

        evidence = str(item.get("evidence", "")).strip()
        status = (
            "supported"
            if evidence and evidence.casefold() in resume_text.casefold()
            else "needs_review"
        )

        skills.append({
            "name": name,
            "category": category,
            "evidence": evidence,
            "validation_status": status
        })

    candidate = raw.get("candidate")
    if not isinstance(candidate, dict):
        candidate = {}

    education = raw.get("education", [])
    projects = raw.get("projects", [])
    warnings = raw.get("extraction_warnings", [])

    if not isinstance(education, list):
        education = []
    if not isinstance(projects, list):
        projects = []
    if not isinstance(warnings, list):
        warnings = []

    return {
        "candidate": {
            "name": candidate.get("name"),
            "email": candidate.get("email"),
            "phone": candidate.get("phone")
        },
        "education": education,
        "projects": projects,
        "skills": skills,
        "extraction_warnings": warnings
    }


@router.post("/analyze", response_model=ResumeResult)
async def analyze_resume(file: UploadFile = File(...)):
    filename = file.filename or ""

    if not filename.lower().endswith((".pdf", ".docx")):
        raise HTTPException(
            status_code=415,
            detail="Upload a PDF or DOCX resume."
        )

    content = await file.read()

    if not content:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    if len(content) > MAX_FILE_BYTES:
        raise HTTPException(
            status_code=413,
            detail="Resume exceeds the 12 MB upload limit."
        )

    resume_text = await extract_text(filename, content)

    if len(resume_text.strip()) < 40:
        raise HTTPException(
            status_code=422,
            detail="No readable text was found. This may be a scanned PDF. Use OCR or upload a text-based PDF/DOCX."
        )

    api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:
        raise HTTPException(
            status_code=503,
            detail="Resume extraction requires GEMINI_API_KEY in the backend environment. No analysis was generated."
        )

    model = os.getenv("GEMINI_MODEL", "gemini-2.5-flash")

    system_prompt = """
Extract only information supported by the supplied resume text.
Return a valid JSON object with these keys:
{
  "candidate": {"name": null, "email": null, "phone": null},
  "education": [{"degree": "", "institution": "", "graduation_year": null}],
  "projects": [{"name": "", "description": "", "technologies": []}],
  "skills": [{"name": "", "category": "OTHER", "evidence": ""}],
  "extraction_warnings": []
}
Allowed skill categories: COD, DSA, OOD, APTI, COMM, AI, CLOUD, SQL, SWE, SYSD, NETW, OS, OTHER.
Use only supported information. Do not invent skills, education, projects, or personal details.
For every skill, provide a short exact quote copied from the resume as evidence.
Use null or empty arrays when information is absent.
Return JSON only without Markdown.
"""

    payload = {
        "model": model,
        "temperature": 0,
        "response_format": {"type": "json_object"},
        "messages": [
            {"role": "system", "content": system_prompt},
            {
                "role": "user",
                "content": "Extract the resume data from this text:\n\n" + resume_text[:50000]
            }
        ]
    }

    gemini_url = (
        "https://generativelanguage.googleapis.com/"
        "v1beta/openai/chat/completions"
    )

    try:
        async with httpx.AsyncClient(timeout=60) as client:
            response = await client.post(
                gemini_url,
                headers={
                    "Authorization": f"Bearer {api_key}",
                    "Content-Type": "application/json"
                },
                json=payload
            )

        if response.status_code >= 400:
            print("Gemini API error:", response.status_code, response.text[:1000])
            raise HTTPException(
                status_code=502,
                detail="Gemini rejected the request. Check GEMINI_API_KEY, GEMINI_MODEL, API access, and backend logs."
            )

        body = response.json()
        content_text = body["choices"][0]["message"]["content"].strip()

        if content_text.startswith("```"):
            content_text = re.sub(
                r"^```(?:json)?\s*|\s*```$",
                "",
                content_text,
                flags=re.IGNORECASE
            ).strip()

        raw = json.loads(content_text)

    except HTTPException:
        raise
    except (httpx.HTTPError, KeyError, IndexError, ValueError) as exc:
        print("Resume analysis error:", type(exc).__name__, str(exc)[:500])
        raise HTTPException(
            status_code=502,
            detail="Resume analysis failed. Check backend logs, Gemini configuration, and model availability."
        ) from exc

    return validate_result(raw, resume_text)


@router.post("/save")
async def save_resume(
    data: str = Form(...),  # JSON string from frontend
    user_id: str = Form(...),
    file_name: str = Form(...),
    file_path: str = Form(""),
    file_type: str = Form(""),
    file_size: int = Form(0),
    linkedin_url: str = Form(""),
    github_url: str = Form(""),
    portfolio_url: str = Form(""),
    professional_summary: str = Form(""),
    total_experience_years: float = Form(0.0),
):
    # Parse the JSON string
    try:
        data_dict = json.loads(data)
    except json.JSONDecodeError as e:
        raise HTTPException(status_code=422, detail=f"Invalid JSON in 'data' field: {e}")

    # Validate against ResumeResult model
    try:
        resume_data = ResumeResult(**data_dict)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Invalid resume data structure: {e}")

    supabase_url = os.getenv("SUPABASE_URL")
    service_key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

    if not supabase_url or not service_key:
        raise HTTPException(
            status_code=503,
            detail="Saving requires SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY in the backend environment."
        )

    table = os.getenv("SUPABASE_RESUME_TABLE", "resumes")

    if not table or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", table):
        raise HTTPException(
            status_code=503,
            detail="Set SUPABASE_RESUME_TABLE to a valid existing table name after confirming its columns. No data was saved."
        )

    # Build parsed_data JSONB with all extracted information
    parsed_data = {
        "education": [item.model_dump() for item in resume_data.education],
        "projects": [item.model_dump() for item in resume_data.projects],
        "skills": [item.model_dump() for item in resume_data.skills],
        "extraction_warnings": resume_data.extraction_warnings,
        "candidate": {
            "name": resume_data.candidate.name,
            "email": resume_data.candidate.email,
            "phone": resume_data.candidate.phone,
        }
    }

    row = {
        "user_id": user_id,
        "file_name": file_name,
        "file_path": file_path,
        "file_type": file_type,
        "file_size": file_size,
        "full_name": resume_data.candidate.name,
        "email": resume_data.candidate.email,
        "phone": resume_data.candidate.phone,
        "location": None,  # Not extracted by current parser
        "linkedin_url": linkedin_url or None,
        "github_url": github_url or None,
        "portfolio_url": portfolio_url or None,
        "professional_summary": professional_summary or None,
        "total_experience_years": total_experience_years,
        "parsing_status": "completed",
        "parsing_error": None,
        "parsed_data": parsed_data,
    }

    try:
        async with httpx.AsyncClient(timeout=30) as client:
            response = await client.post(
                f"{supabase_url.rstrip('/')}/rest/v1/{table}",
                headers={
                    "apikey": service_key,
                    "Authorization": f"Bearer {service_key}",
                    "Content-Type": "application/json",
                    "Prefer": "return=representation"
                },
                json=row
            )

        if response.status_code >= 400:
            print("Supabase save error:", response.status_code, response.text[:1000])
            raise HTTPException(
                status_code=502,
                detail=f"Supabase rejected the save: {response.text[:500]}"
            )

    except HTTPException:
        raise
    except httpx.HTTPError as exc:
        raise HTTPException(
            status_code=502,
            detail="Could not connect to Supabase while saving the resume."
        ) from exc

    return {"ok": True, "saved_id": response.json()[0]["id"] if response.json() else None}